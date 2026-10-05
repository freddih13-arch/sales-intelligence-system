"""
Tests para módulo cockpit.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from sales_intel.cockpit.generar_cockpit import (
    generar_prospectos_sinteticos,
    crear_hoja_cockpit,
    crear_hoja_datos_fuente,
    crear_hoja_listas,
    crear_hoja_exportar_csv,
    VOCABULARIOS,
    generar_cockpit_demo,
)


def test_vocabularios_completos():
    """Verifica que todos los vocabularios tengan valores."""
    for nombre, valores in VOCABULARIOS.items():
        assert isinstance(valores, list), f"{nombre} debe ser lista"
        assert len(valores) > 0, f"{nombre} no puede estar vacío"
        for v in valores:
            assert isinstance(v, str), f"Valores en {nombre} deben ser strings"


def test_generar_prospectos_sinteticos():
    """Test generación de prospectos sintéticos."""
    prospectos = generar_prospectos_sinteticos(5)
    assert len(prospectos) == 5

    for p in prospectos:
        assert "entidad_dedup_id" in p
        assert "razon_social" in p
        assert "sector_vertical" in p
        assert "score_prioridad_comercial" in p
        assert "prioridad" in p
        assert p["prioridad"] in ["ALTA", "MEDIA", "BAJA"]
        assert 0 <= p["score_prioridad_comercial"] <= 100


def test_generar_cockpit_demo():
    """Test generación completa de cockpit demo."""
    import tempfile
    with tempfile.TemporaryDirectory() as tmpdir:
        salida = Path(tmpdir) / "cockpit_test.xlsx"
        resultado = generar_cockpit_demo(str(salida), n_prospectos=3)

        assert Path(resultado).exists()
        assert Path(resultado).stat().st_size > 0


def test_crear_hoja_cockpit():
    """Test creación de hoja cockpit con estructura correcta."""
    from openpyxl import Workbook

    prospectos = generar_prospectos_sinteticos(2)
    wb = Workbook()
    crear_hoja_cockpit(wb, prospectos)

    ws = wb["COCKPIT"]
    assert ws.title == "COCKPIT"

    # Verificar que existe selector en C3
    assert ws["C3"].value is not None
    assert "PROSPECTO" in str(ws["B3"].value).upper()


def test_hojas_auxiliares():
    """Test creación de hojas auxiliares."""
    from openpyxl import Workbook

    prospectos = generar_prospectos_sinteticos(2)
    wb = Workbook()
    crear_hoja_datos_fuente(wb, prospectos)
    assert "DATOS_FUENTE" in wb.sheetnames
    assert wb["DATOS_FUENTE"].sheet_state == "hidden"

    crear_hoja_listas(wb)
    assert "LISTAS" in wb.sheetnames
    assert wb["LISTAS"].sheet_state == "hidden"

    crear_hoja_exportar_csv(wb, prospectos)
    assert "EXPORTAR_CSV" in wb.sheetnames


def test_vocabularios_en_listas():
    """Verifica que vocabularios se escriban en hoja LISTAS."""
    from openpyxl import Workbook

    prospectos = generar_prospectos_sinteticos(1)
    wb = Workbook()
    crear_hoja_listas(wb)

    ws = wb["LISTAS"]
    # Verificar que al menos un vocabulario se escribió
    valores = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert len(valores) > 0
    assert "resultado_contacto" in valores or "etapa_actual" in valores


def test_generar_prospectos_reproducible():
    """Test que la generación sea reproducible con seed fijo."""
    p1 = generar_prospectos_sinteticos(5)
    p2 = generar_prospectos_sinteticos(5)
    # Mismo seed por defecto = mismos resultados
    for i in range(5):
        assert p1[i]["entidad_dedup_id"] == p2[i]["entidad_dedup_id"]
        assert p1[i]["razon_social"] == p2[i]["razon_social"]
        assert p1[i]["score_prioridad_comercial"] == p2[i]["score_prioridad_comercial"]


if __name__ == "__main__":
    import sys
    fallos = 0
    for nombre, funcion in list(globals().items()):
        if nombre.startswith("test_") and callable(funcion):
            try:
                funcion()
                print(f"OK   {nombre}")
            except AssertionError as e:
                fallos += 1
                print(f"FAIL {nombre}: {e}")
            except Exception as e:
                fallos += 1
                print(f"ERROR {nombre}: {e}")
    raise SystemExit(fallos)